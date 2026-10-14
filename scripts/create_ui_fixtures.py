"""Create fictional documents for repeatable browser and packaged-parser checks."""

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_bytes(blank_page=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 50 790 Td (Paziente: Mario Rossi - PA 135/85) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    if blank_page:
        writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def main():
    folder = Path(__file__).resolve().parents[1] / "build/upload-fixtures"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "note.pdf").write_bytes(pdf_bytes())
    (folder / "partial.pdf").write_bytes(pdf_bytes(blank_page=True))
    word = Document()
    word.add_paragraph("Paziente: Mario Rossi")
    word.add_paragraph("PA 135/85 mmHg")
    word.save(folder / "note.docx")
    print("Fictional PDF/DOCX browser fixtures created.")


if __name__ == "__main__":
    main()

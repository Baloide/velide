import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZIP_LZMA, ZipFile

from pypdf import get_configuration

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.main import app
from app.services.text_ingestion import MAX_FILE_BYTES, IngestionError, extract_text
from docx import Document
from docx.enum.section import WD_SECTION_START
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_fixture(text="Paziente: Mario Rossi", blank_page=False, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 790 Td (" + text.encode("ascii") + b") Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    if blank_page:
        writer.add_blank_page(width=595, height=842)
    if encrypted:
        writer.encrypt("test-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_fixture():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Dott. Anna Bianchi"
    document.add_paragraph("Paziente: Mario Rossi")
    identity = document.add_table(rows=1, cols=2)
    identity.cell(0, 0).text = "Paziente:"
    identity.cell(0, 1).text = "Maria Verdi"
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Email: mario.rossi@example.com"
    table.cell(0, 1).text = "PA 135/85 mmHg"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "FC 88 bpm"
    document.add_paragraph("Visita: 14 ottobre 2026")
    document.sections[0].footer.paragraphs[0].text = "Cartella Clinica: CC-2026-1042"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def upload(self, data, format):
        return self.client.post(
            "/api/ingest/file",
            params={"format": format},
            content=data,
            headers={"Content-Type": "application/octet-stream"},
        )

    def test_docx_body_tables_nested_headers_footers(self):
        result = extract_text(docx_fixture(), "docx")
        for value in [
            "Mario Rossi",
            "Anna Bianchi",
            "mario.rossi@example.com",
            "135/85",
            "88 bpm",
            "CC-2026-1042",
        ]:
            self.assertIn(value, result.text)
        self.assertEqual(result.text.count("Anna Bianchi"), 1)
        self.assertLess(
            result.text.index("Mario Rossi"), result.text.index("mario.rossi@example.com")
        )
        self.assertLess(result.text.index("88 bpm"), result.text.index("14 ottobre"))
        self.assertIn("docxReview", result.warnings)
        self.assertEqual(result.characters, len(result.text))
        self.assertIn("Paziente:\tMaria Verdi", result.text)

    def test_pdf_and_partial_pdf(self):
        result = extract_text(pdf_fixture(blank_page=True), "pdf")
        self.assertIn("Mario Rossi", result.text)
        self.assertEqual(result.pages, 2)
        self.assertEqual(result.pages_without_text, [2])
        self.assertIn("partialPdf", result.warnings)

    def test_names_in_imported_roster_and_author_list(self):
        document = Document()
        document.add_paragraph("Problem Statement")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Nome"
        table.rows[0].cells[1].text = "Matricola"
        for name, identifier in [
            ("La Magna Davide", "0512121375"),
            ("Ciccarelli Raffaele", "0512121405"),
            ("Manzolillo Matteo", "0512121666"),
        ]:
            row = table.add_row()
            row.cells[0].text = name
            row.cells[1].text = identifier
        document.add_paragraph(
            "Scritto da: Ciccarelli Raffaele, La Magna Davide, Manzolillo Matteo"
        )
        document.add_paragraph("Problem Statement")
        stream = BytesIO()
        document.save(stream)
        extracted = self.upload(stream.getvalue(), "docx").json()["text"]
        response = self.client.post(
            "/api/redact/text", json={"text": extracted, "options": {"language": "it"}}
        ).json()
        for name in ["La Magna Davide", "Ciccarelli Raffaele", "Manzolillo Matteo"]:
            self.assertNotIn(name, response["redacted_text"])
        self.assertEqual(len([e for e in response["entities"] if e["entity_type"] == "ID"]), 3)
        self.assertIn("Problem Statement", response["redacted_text"])

    def test_text_formats_and_encodings(self):
        for format in ["txt", "md"]:
            for encoding in ["utf-8-sig", "utf-16", "cp1252"]:
                result = extract_text("Paziente: André Rossi\r\nVisita".encode(encoding), format)
                self.assertEqual(result.text, "Paziente: André Rossi\nVisita")

    def test_api_extract_then_redact_each_format(self):
        for format, data in [
            ("pdf", pdf_fixture()),
            ("docx", docx_fixture()),
            ("txt", b"Paziente: Mario Rossi"),
            ("md", b"Paziente: Mario Rossi"),
        ]:
            with self.subTest(format=format):
                response = self.upload(data, format)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers["cache-control"], "no-store")
                self.assertEqual(response.json()["format"], format)
                redaction = self.client.post(
                    "/api/redact/text",
                    json={"text": response.json()["text"], "options": {"language": "it"}},
                )
                self.assertNotIn("Mario Rossi", redaction.json()["redacted_text"])
                self.assertNotIn("Anna Bianchi", redaction.json()["redacted_text"])
                self.assertNotIn("Maria Verdi", redaction.json()["redacted_text"])

    def test_empty_scanned_encrypted_corrupt_and_binary(self):
        blank = PdfWriter()
        blank.add_blank_page(width=595, height=842)
        stream = BytesIO()
        blank.write(stream)
        cases = [
            ("pdf", stream.getvalue(), "pdfNoText"),
            ("pdf", pdf_fixture(encrypted=True), "encryptedPdf"),
            ("pdf", b"not pdf", "unreadableFile"),
            ("docx", b"not docx", "unreadableFile"),
            ("txt", b"", "emptyFile"),
            ("md", b"   ", "emptyFile"),
            ("txt", b"\x00\x01binary", "textEncoding"),
        ]
        for format, data, code in cases:
            with self.subTest(code=code):
                response = self.upload(data, format)
                self.assertEqual(response.status_code, 422)
                self.assertEqual(response.json()["detail"], {"code": code})

    def test_limits(self):
        response = self.upload(b"x" * (MAX_FILE_BYTES + 1), "txt")
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["detail"]["code"], "fileTooLarge")
        response = self.upload(b"x" * 200001, "txt")
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["detail"]["code"], "textTooLong")
        with patch("app.services.text_ingestion.MAX_PDF_PAGES", 1):
            response = self.upload(pdf_fixture(blank_page=True), "pdf")
            self.assertEqual(response.json()["detail"]["code"], "tooManyPages")
        with patch("app.services.text_ingestion.MAX_DOCX_EXPANDED_BYTES", 100):
            response = self.upload(docx_fixture(), "docx")
            self.assertEqual(response.json()["detail"]["code"], "documentTooComplex")
        self.assertEqual(self.upload(b"doc", "doc").status_code, 422)

    def test_pdf_expansion_limit_and_context_restoration(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=595, height=842)
        stream = DecodedStreamObject()
        stream.set_data(b" " * 4096)
        page[NameObject("/Contents")] = writer._add_object(stream.flate_encode())
        data = BytesIO()
        writer.write(data)
        previous = get_configuration()
        with patch("app.services.text_ingestion.MAX_PDF_STREAM_BYTES", 128):
            with self.assertRaises(IngestionError) as failure:
                extract_text(data.getvalue(), "pdf")
            self.assertEqual(failure.exception.code, "documentTooComplex")
            self.assertEqual(get_configuration(), previous)
            response = self.upload(data.getvalue(), "pdf")
            self.assertEqual(response.status_code, 413)
            self.assertEqual(response.json()["detail"]["code"], "documentTooComplex")
        self.assertEqual(get_configuration(), previous)
        self.assertEqual(self.upload(pdf_fixture(), "pdf").status_code, 200)

    def test_docx_component_count_and_compression(self):
        stream = BytesIO()
        with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "<document/>")
            for number in range(1000):
                archive.writestr(f"part-{number}", "")
        response = self.upload(stream.getvalue(), "docx")
        self.assertEqual(response.status_code, 413)
        stream = BytesIO()
        with ZipFile(stream, "w", ZIP_LZMA) as archive:
            archive.writestr("word/document.xml", "<document/>")
        self.assertEqual(self.upload(stream.getvalue(), "docx").status_code, 422)

    def test_docx_nested_table_depth_guard(self):
        document = Document()
        container = document
        for _ in range(22):
            container = container.add_table(rows=1, cols=1).cell(0, 0)
        container.text = "Paziente: Mario Rossi"
        stream = BytesIO()
        document.save(stream)
        response = self.upload(stream.getvalue(), "docx")
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["detail"]["code"], "documentTooComplex")


if __name__ == "__main__":
    unittest.main()
